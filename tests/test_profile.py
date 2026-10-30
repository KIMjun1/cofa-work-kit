#!/usr/bin/env python3
"""profile_data.py 검증. 저장소 루트에서: python3 tests/test_profile.py"""
import subprocess
import sys
import tempfile
import time
from pathlib import Path

SCRIPT = Path(__file__).resolve().parent.parent / "tools" / "profile_data.py"


def make_fixtures(d: Path):
    import openpyxl
    import pandas as pd
    pd.DataFrame({"라인": ["A", "B", "C"], "수율": [95.1, None, 93.2]}).to_csv(
        d / "kor_cp949.csv", index=False, encoding="cp949")
    pd.DataFrame({"x": range(30000)}).to_csv(d / "rows30000.csv", index=False)
    (d / "semicolon.csv").write_text("line;qty;yield\nA;10;0,95\nB;20;0,93\nC;30;0,91\n", encoding="utf-8")
    pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", "z"]}).to_csv(d / "utf16.csv", index=False, encoding="utf-16")
    (d / "thousands.csv").write_text('id,amount\n1,"1,234"\n2,"2,500"\n3,"10,000"\n4,-\n', encoding="utf-8")
    with pd.ExcelWriter(d / "two_sheets.xlsx") as w:
        pd.DataFrame({"a": range(1000), "b": [None] * 500 + list(range(500))}).to_excel(w, sheet_name="S1", index=False)
        pd.DataFrame({"c": ["x", "y"]}).to_excel(w, sheet_name="S2", index=False)
    # 서식만 있는 빈 꼬리 행: 실제 100행인데 시트 범위는 5000행
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["k", "v"])
    for i in range(100):
        ws.append([i, i * 2])
    ws.cell(row=5000, column=1).number_format = "0.00"
    wb.save(d / "tail_format.xlsx")
    # 범위 정보가 없는 파일 (write_only)
    wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet("W")
    ws.append(["k"])
    for i in range(250):
        ws.append([i])
    wb.save(d / "write_only.xlsx")
    (d / "corrupted.xlsx").write_bytes(b"not an excel file")
    (d / "notes.txt").write_text("memo", encoding="utf-8")
    tools_dir = d / "tools"; tools_dir.mkdir()
    (tools_dir / "lookup.csv").write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    try:
        from PIL import Image
        Image.new("RGB", (64, 48)).save(d / "wafer.tif")
    except ImportError:
        pass


def main():
    with tempfile.TemporaryDirectory() as data, tempfile.TemporaryDirectory() as cwd:
        make_fixtures(Path(data))
        t0 = time.time()
        r = subprocess.run([sys.executable, str(SCRIPT), data], cwd=cwd, capture_output=True, text=True)
        elapsed = time.time() - t0
        assert r.returncode == 0, r.stderr
        out = (Path(cwd) / "DATA_PROFILE.md").read_text(encoding="utf-8")
        checks = {
            "30초 이내": elapsed < 30,
            "cp949 CSV": "인코딩: cp949" in out,
            "CSV 전체 행 수": "30000행 (전체" in out,
            "세미콜론 구분자": "구분자: 세미콜론" in out and "semicolon.csv" in out,
            "UTF-16 CSV": "인코딩: utf-16" in out,
            "천 단위 쉼표 경고": "열 'amount'은 숫자로 보이나" in out,
            "엑셀 실제 행 수": "[S1] 1000행 (전체" in out,
            "엑셀 결측 비율": "50.0%" in out,
            "두 번째 시트": "[S2] 2행" in out,
            "빈 꼬리 행 무시": "[Sheet] 100행 (전체" in out,
            "범위 정보 없는 엑셀": "[W] 250행 (전체" in out,
            "손상 파일 격리": "corrupted.xlsx (처리 실패" in out,
            "기타 파일 목록": "notes.txt" in out,
            "과제의 tools/ 폴더 포함": "lookup.csv" in out,
            "tif 이미지": "wafer.tif (TIFF" in out or "Pillow 미설치" in out,
        }
        for k, ok in checks.items():
            print(("[PASS] " if ok else "[FAIL] ") + k)
        print(f"소요 {elapsed:.1f}초")
        if not all(checks.values()):
            print(out)
            sys.exit(1)


if __name__ == "__main__":
    main()
